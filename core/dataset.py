"""
Tabular data ingestion for uploaded clinical files.

Reads CSV and XLSX (via openpyxl, if installed) into a uniform
{"columns": [...], "rows": [ {col: value} ]} shape, then derives the summary
statistics the generated documents cite.
"""

from __future__ import annotations

import csv
import hashlib
import os
from typing import Any

TABULAR_EXTENSIONS = {".csv", ".xlsx", ".xls"}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".doc", ".txt", ".potx", ".pptx"}
ALLOWED_EXTENSIONS = TABULAR_EXTENSIONS | DOCUMENT_EXTENSIONS

MAX_PREVIEW_ROWS = 25


def extension(filename: str) -> str:
    return os.path.splitext(filename)[1].lower()


def is_allowed(filename: str) -> bool:
    return extension(filename) in ALLOWED_EXTENSIONS


def is_tabular(filename: str) -> bool:
    return extension(filename) in TABULAR_EXTENSIONS


def checksum(path: str) -> str:
    """SHA-256 of a file on disk, for the reproducibility record."""
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(65536), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_table(path: str) -> dict[str, Any]:
    """
    Read a CSV/XLSX file into {"columns", "rows", "row_count", "error"}.

    Never raises: an unreadable file comes back with an `error` string so the UI
    can show the problem next to the file instead of 500-ing the upload.
    """
    ext = extension(path)
    try:
        if ext == ".csv":
            return _read_csv(path)
        if ext in (".xlsx", ".xls"):
            return _read_xlsx(path)
    except Exception as exc:
        return {"columns": [], "rows": [], "row_count": 0, "error": f"{type(exc).__name__}: {exc}"}
    return {"columns": [], "rows": [], "row_count": 0, "error": "Unsupported file type"}


def _read_csv(path: str) -> dict[str, Any]:
    with open(path, "r", newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        rows = [dict(row) for row in reader]
        columns = list(reader.fieldnames or [])
    return {"columns": columns, "rows": rows, "row_count": len(rows), "error": None}


def _read_xlsx(path: str) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {
            "columns": [],
            "rows": [],
            "row_count": 0,
            "error": "XLSX support needs openpyxl — run: pip install -r requirements.txt",
        }

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    grid = sheet.iter_rows(values_only=True)

    try:
        header = next(grid)
    except StopIteration:
        workbook.close()
        return {"columns": [], "rows": [], "row_count": 0, "error": "Worksheet is empty"}

    columns = [str(cell) if cell is not None else f"column_{i}" for i, cell in enumerate(header)]
    rows = []
    for record in grid:
        if all(cell is None for cell in record):
            continue
        rows.append({
            columns[i]: ("" if value is None else value)
            for i, value in enumerate(record)
            if i < len(columns)
        })
    workbook.close()
    return {"columns": columns, "rows": rows, "row_count": len(rows), "error": None}


# --------------------------------------------------------------------------
# Derived statistics
# --------------------------------------------------------------------------


def _numeric(rows: list[dict], column: str) -> list[float]:
    values = []
    for row in rows:
        raw = row.get(column)
        if raw in (None, ""):
            continue
        try:
            values.append(float(str(raw).strip().rstrip("%")))
        except ValueError:
            continue
    return values


def _mean(values: list[float]) -> float | None:
    return round(sum(values) / len(values), 1) if values else None


def summarise_study(rows: list[dict]) -> dict[str, Any]:
    """Cohort-level statistics for the pivotal-study table."""
    if not rows:
        return {}

    ages = _numeric(rows, "Age")
    mards = _numeric(rows, "MARD")
    wear = _numeric(rows, "Sensor_Wear_Days")
    hba1c = _numeric(rows, "Baseline_HbA1c")

    types = [str(row.get("Diabetes_Type", "")).strip() for row in rows]
    type_1 = sum(1 for value in types if value in ("1", "Type 1", "T1D"))

    genders = [str(row.get("Gender", "")).strip().upper()[:1] for row in rows]
    female = sum(1 for value in genders if value == "F")

    return {
        "patients": len(rows),
        "mean_age": _mean(ages),
        "mean_mard": _mean(mards),
        "mean_wear_days": _mean(wear),
        "mean_hba1c": _mean(hba1c),
        "type_1_count": type_1,
        "type_1_pct": round(100 * type_1 / len(rows)) if rows else 0,
        "female_pct": round(100 * female / len(rows)) if rows else 0,
    }


def summarise_safety(rows: list[dict]) -> dict[str, Any]:
    """Adverse-event counts by type and severity."""
    if not rows:
        return {}

    by_type: dict[str, int] = {}
    by_severity: dict[str, int] = {}
    for row in rows:
        event = str(row.get("Event_Type", "Unspecified")).strip()
        severity = str(row.get("Severity", "Unspecified")).strip()
        by_type[event] = by_type.get(event, 0) + 1
        by_severity[severity] = by_severity.get(severity, 0) + 1

    resolved = sum(
        1 for row in rows if str(row.get("Resolved", "")).strip().lower() in ("yes", "true", "1")
    )
    serious = by_severity.get("Severe", 0) + by_severity.get("Serious", 0)

    return {
        "total": len(rows),
        "serious": serious,
        "resolved": resolved,
        "resolution_pct": round(100 * resolved / len(rows)) if rows else 0,
        "by_type": dict(sorted(by_type.items(), key=lambda item: -item[1])),
        "by_severity": by_severity,
    }


def summarise_efficacy(rows: list[dict]) -> dict[str, Any]:
    """Endpoint table keyed by metric name, keeping CIs and p-values intact."""
    endpoints = {}
    for row in rows:
        metric = str(row.get("Metric", "")).strip()
        if not metric:
            continue
        endpoints[metric] = {
            "value": row.get("Value"),
            "unit": row.get("Unit", ""),
            "ci_low": row.get("95_CI_Lower"),
            "ci_high": row.get("95_CI_Upper"),
            "p_value": row.get("P_Value", ""),
            "significance": row.get("Clinical_Significance", ""),
        }
    return {"endpoints": endpoints, "count": len(endpoints)}


def classify(filename: str, columns: list[str]) -> str:
    """Guess a data category from the filename and header row."""
    haystack = (filename + " " + " ".join(columns)).lower()
    if any(token in haystack for token in ("adverse", "safety", "event_type", "severity")):
        return "safety"
    if any(token in haystack for token in ("efficacy", "endpoint", "p_value", "metric")):
        return "efficacy"
    if any(token in haystack for token in ("demographic", "baseline", "age", "gender")):
        return "demographics"
    if any(token in haystack for token in ("tender", "pricing", "cost", "budget")):
        return "economics"
    return "clinical"
